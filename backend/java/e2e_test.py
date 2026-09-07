# -*- coding: utf-8 -*-
"""Java 后端端到端测试：构造模板/Excel → 上传 → 保存配置 → 制函 → 断言输出"""
import json, os, io, uuid, urllib.request, urllib.error

BASE = r"e:\13_dingdian\03_demo\04_Letter_Gen"
JAVA_URL = "http://127.0.0.1:5002"
TMP = os.path.join(BASE, "backend", "java", "_tmp_test")
os.makedirs(TMP, exist_ok=True)

passed, failed = [], []
def check(name, cond, detail=""):
    (passed if cond else failed).append(name)
    print(("PASS" if cond else "FAIL") + f": {name}" + (f" ({detail})" if detail and not cond else ""))

# ---- 构造测试 Word 模板（含占位段）----
from docx import Document
tpl_path = os.path.join(TMP, "tpl_java_test.docx")
doc = Document()
doc.add_paragraph("致：被询证单位")
doc.add_paragraph("【Sheet「银行存款」表格将在此处展示】")
doc.add_paragraph("【Sheet「无数据表」表格将在此处展示】")
doc.add_paragraph("函证编号：【无】")
doc.save(tpl_path)

# ---- 构造测试 Excel（1 个分组 Sheet + 1 个无数据分组 Sheet）----
from openpyxl import Workbook
xlsx_path = os.path.join(TMP, "data_java_test.xlsx")
wb = Workbook()
ws = wb.active; ws.title = "银行存款"
ws.append(["被审计单位", "被询证单位", "科目", "金额"])
ws.append(["顶点软件", "中信证券股份有限公司", "银行存款", 1500000])
ws.append(["顶点软件", "国泰君安证券股份有限公司", "银行存款", 2300000.55])
ws2 = wb.create_sheet("无数据表")
ws2.append(["被审计单位", "被询证单位", "事项"])
wb.save(xlsx_path)

# ---- multipart 工具 ----
def multipart(fields, files):
    boundary = uuid.uuid4().hex
    body = io.BytesIO()
    for k, v in fields.items():
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode("utf-8"))
    for k, (fname, fpath, ctype) in files.items():
        body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{fname}\"\r\n"
                   f"Content-Type: {ctype}\r\n\r\n".encode("utf-8"))
        body.write(open(fpath, "rb").read())
        body.write(b"\r\n")
    body.write(f"--{boundary}--\r\n".encode("utf-8"))
    return body.getvalue(), f"multipart/form-data; boundary={boundary}"

def post_json(path, obj):
    req = urllib.request.Request(JAVA_URL + path, data=json.dumps(obj, ensure_ascii=False).encode("utf-8"),
                                 headers={"Content-Type": "application/json; charset=utf-8"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            return r.status, json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode("utf-8"))

# ---- 1. 上传模板 ----
body, ctype = multipart({}, {"file": ("tpl_java_test.docx", tpl_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
req = urllib.request.Request(JAVA_URL + "/api/upload_template", data=body, headers={"Content-Type": ctype}, method="POST")
with urllib.request.urlopen(req, timeout=30) as r:
    d = json.loads(r.read().decode("utf-8"))
check("upload_template 200 + para/table 信息", "path" in d and d.get("table_count") == 0, str(d)[:120])

# ---- 2. 上传批量 Excel ----
body, ctype = multipart({}, {"file": ("data_java_test.xlsx", xlsx_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
req = urllib.request.Request(JAVA_URL + "/api/upload_excel", data=body, headers={"Content-Type": ctype}, method="POST")
with urllib.request.urlopen(req, timeout=30) as r:
    d = json.loads(r.read().decode("utf-8"))
check("upload_excel batch_mode", d.get("batch_mode") is True)
check("upload_excel groups=2", len(d.get("groups", [])) == 2, str(d.get("groups")))
check("upload_excel match=2", len([m for m in d.get("match_results", []) if m.get("matched")]) == 2)
check("upload_excel sheet 结构", d.get("sheets", [{}])[0].get("header", [])[0] == "被审计单位")
excel_uploaded_path = d["path"]

# ---- 3. 保存模板配置（含样式 + 无数据 Sheet 结构）----
code, d = post_json("/api/templates/save", {
    "name": "JAVA测试模板",
    "excel_path": excel_uploaded_path,
    "style": {"font_size": 12, "font_name": "仿宋", "h_align": "center", "v_align": "center",
              "row_height_mode": "atLeast", "row_height_pt": 22,
              "col_widths_mode": "ratio", "col_widths": {"银行存款": "1:2"}},
})
check("templates/save 200", code == 200, str(d))
template_id = d.get("id")
check("save style 落盘", (d.get("style") or {}).get("font_size") == 12.0)
check("save sheets 保留", len(d.get("sheets", [])) == 2)

# ---- 4. 编辑模式再保存（不带 excel_path，样式修改）----
code, d = post_json("/api/templates/save", {
    "id": template_id, "name": "JAVA测试模板-改名",
    "style": {"font_size": 11, "font_name": "", "h_align": "left", "v_align": "top",
              "row_height_mode": "auto", "row_height_pt": 20,
              "col_widths_mode": "auto", "col_widths": {}},
})
check("编辑模式无 excel_path 保存 200", code == 200, str(d))
check("编辑改名生效", d.get("name") == "JAVA测试模板-改名")

# ---- 5. 制函 ----
code, d = post_json("/api/generate", {"template_id": template_id, "excel_path": excel_uploaded_path})
check("generate 200", code == 200, str(d)[:200])
check("generate count=2（无数据 Sheet 不产函证，但两组合均登记匹配）", d.get("count") == 2, str(d.get("count")))
check("generate sheet_missing 为空（无数据表存在但无数据，走占位段删除而非 missing）", d.get("sheet_missing") == [], str(d.get("sheet_missing")))
check("generate unmatched 空", d.get("unmatched") == [])

# ---- 6. 下载 ZIP 并用 python-docx 断言内容/样式 ----
zip_url = JAVA_URL + d["download_url"]
with urllib.request.urlopen(zip_url, timeout=30) as r:
    zip_bytes = r.read()
check("zip 下载非空", len(zip_bytes) > 2000)
import zipfile
zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
names = zf.namelist()
check("zip 含 2 封", len(names) == 2, str(names))
doc_bytes = zf.read(names[0])

from docx import Document as PyDocx
dd = Document(io.BytesIO(doc_bytes))
all_text = "\n".join(p.text for p in dd.paragraphs)
check("占位段零残留", "【Sheet" not in all_text and "将在此处展示" not in all_text)
check("函证正文保留", "致：被询证单位" in all_text)
check("注入 1 个表格（无数据表已删）", len(dd.tables) == 1, str(len(dd.tables)))
t = dd.tables[0]
check("表头正确", [c.text for c in t.rows[0].cells] == ["科目", "金额"], str([c.text for c in t.rows[0].cells]))
check("数据正确（数值无小数尾巴）", t.rows[1].cells[1].text == "1500000", t.rows[1].cells[1].text)
r0 = t.rows[0].cells[0].paragraphs[0].runs[0]
check("样式：字号 11 生效", r0.font.size is not None and abs(r0.font.size.pt - 11) < 0.01, str(r0.font.size))
check("样式：表头加粗", r0.bold is True)

# ---- 7. 编辑模式样式回显 ----
with urllib.request.urlopen(JAVA_URL + "/api/templates", timeout=10) as r:
    items = json.loads(r.read().decode("utf-8"))["templates"]
it = next(x for x in items if x["id"] == template_id)
check("列表接口回显 style", (it.get("style") or {}).get("col_widths_mode") == "auto")

# ---- 清理 ----
req = urllib.request.Request(JAVA_URL + f"/api/templates/{template_id}", method="DELETE")
with urllib.request.urlopen(req, timeout=10) as r:
    check("删除模板 200", json.loads(r.read().decode("utf-8")).get("ok") is True)

print()
print(f"==== RESULT: {len(passed)} passed, {len(failed)} failed ====")
if failed:
    print("FAILED:", failed)
