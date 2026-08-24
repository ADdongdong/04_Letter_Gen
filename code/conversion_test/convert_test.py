# -*- coding: utf-8 -*-
"""docx→markdown 还原度实验：验证函证模板转换效果"""
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

SRC = r'E:\13_dingdian\02_产品\06_函证系统\01_售前演示数据\函证模板\往来函证模板（销售、采购）.docx'
OUT_DIR = r'e:\13_dingdian\03_demo\04_制函优化\上下文\conversion_test'
os.makedirs(OUT_DIR, exist_ok=True)
MD_OUT = os.path.join(OUT_DIR, 'template_from_docx2markdown.md')

# ---- Part 1: docx2markdown 转换 ----
from docx2markdown import docx_to_markdown
print('=== Part 1: docx2markdown ===')
try:
    docx_to_markdown(SRC, MD_OUT)
    print('转换输出文件:', MD_OUT)
    with open(MD_OUT, encoding='utf-8') as f:
        md = f.read()
    print('Markdown 总长度(字符):', len(md))
    print('Markdown 总行数:', md.count('\n'))
except Exception as e:
    print('转换失败:', type(e).__name__, e)

# ---- Part 2: python-docx 提取原始结构（用于对比） ----
print('\n=== Part 2: python-docx 原始结构 ===')
from docx import Document
doc = Document(SRC)
paras = doc.paragraphs
print('段落总数:', len(paras))
print('表格总数:', len(doc.tables))
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_cell = t.rows[0].cells[0].text[:20] if rows else ''
    print(f'  表{i+1}: {rows}行 x {cols}列 | 首格: {first_cell!r}')

# ---- Part 3: 检查转换后 md 里是否保留表格 ----
print('\n=== Part 3: Markdown 中表格/结构还原检查 ===')
if os.path.exists(MD_OUT):
    with open(MD_OUT, encoding='utf-8') as f:
        md = f.read()
    lines = md.split('\n')
    table_rows = [l for l in lines if l.strip().startswith('|')]
    print('Markdown 中的表格行(以|开头)数量:', len(table_rows))
    # 检查 Output 函数/{{变量}} 是否保留
    import re
    output_count = len(re.findall(r'Output\s*\(', md))
    var_count = len(re.findall(r'\{\{', md))
    print('Output( 出现次数:', output_count)
    print('{{ 变量 出现次数:', var_count)
    # 打印前 60 行预览
    print('\n----- Markdown 预览(前60行) -----')
    print('\n'.join(lines[:60]))
